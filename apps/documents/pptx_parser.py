"""연결용 신규 파서 · 오대웅(DaeUng79) 문서 모듈 검토 대상.

PPTX 제안서를 hwpx_parser와 같은 문서 JSON 모양으로 바꾼다. 결과는 그대로
`hwpx_storage.save_to_sqlite`에 넣을 수 있고, 평가 화면의 DocumentReader와
step3 요구사항 추출(compact_blocks)이 그대로 읽는다.

- 물리 페이지 1장 = 슬라이드 1장(슬라이드 순서). page_number도 슬라이드 번호다.
- 도형은 그룹 안까지 재귀로 돌며 (top, left) 순으로 읽는다. 제목 자리표시자는 항상 먼저 읽는다.
- 텍스트 틀의 문단 하나 = 문단 블록 하나. 표 도형 = 표 블록 하나(병합 셀은 빈 셀로 두고 격자를 유지).
- 그림·차트·슬라이드 노트는 읽지 않고 warnings에 한 번 남긴다.

이 모듈은 import 시점에 서드파티를 불러오지 않는다. python-pptx는 parse_pptx 안에서만 쓴다.
문서 뼈대·블록 생성 함수는 pdf_parser와 같이 쓴다.

예: python apps/documents/pptx_parser.py 제안서.pptx
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

try:
    from .pdf_parser import (add_paragraph_block, add_table_block, document_skeleton,
                             finish_document, normalize_space, page_entry)
except ImportError:  # 단독 실행 (apps/documents가 sys.path에 있을 때)
    from pdf_parser import (add_paragraph_block, add_table_block, document_skeleton,
                            finish_document, normalize_space, page_entry)

PPTX_PARSER_VERSION = "pptx-text-1.0"
PAGE_SOURCE = "pptx_slide"


def _clean_lines(text: str) -> str:
    """python-pptx 글의 줄바꿈(\\v, \\n)을 '\\n'으로 통일하고 각 줄의 공백을 정리한다. 빈 줄은 버린다."""
    lines = [normalize_space(part) for part in text.replace("\v", "\n").split("\n")]
    return "\n".join(line for line in lines if line)


def _shape_key(shape) -> tuple[int, int]:
    """(top, left) 정렬 키. 자리표시자처럼 위치가 없으면 0으로 본다."""
    top = shape.top if shape.top is not None else 0
    left = shape.left if shape.left is not None else 0
    return (int(top), int(left))


def _shape_type(shape):
    """shape_type. python-pptx가 모르는 도형은 NotImplementedError를 내므로 None으로 돌린다."""
    try:
        return shape.shape_type
    except NotImplementedError:
        return None


def _walk_shapes(shapes, group_type):
    """그룹 도형 안까지 (top, left) 순으로 도형을 낸다."""
    for shape in sorted(shapes, key=_shape_key):
        if _shape_type(shape) == group_type:
            yield from _walk_shapes(shape.shapes, group_type)
        else:
            yield shape


def _table_grid(table) -> list[list[str]]:
    """표의 셀 글을 직사각형 격자로 뽑는다. 병합으로 가려진 셀은 빈 글로 둔다."""
    grid: list[list[str]] = []
    for row in table.rows:
        values = []
        for cell in row.cells:
            if getattr(cell, "is_spanned", False):
                values.append("")
            else:
                values.append(_clean_lines(cell.text or ""))
        grid.append(values)
    return grid


def parse_pptx(path: str | Path) -> dict[str, Any]:
    """PPTX에서 슬라이드별 글·표를 뽑아 문서 JSON을 만든다.

    RuntimeError: python-pptx가 없을 때. ValueError: 손상 파일, 글자가 전혀 없는 파일.
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.exc import PackageNotFoundError
    except ImportError as error:
        raise RuntimeError(
            "PPTX 제안서를 읽으려면 python-pptx 패키지가 필요합니다. apps/documents/requirements.txt를 확인하세요."
        ) from error

    import io

    path = Path(path)
    data = path.read_bytes()
    document = document_skeleton(path, data, PPTX_PARSER_VERSION, "pptx")
    warnings = document["warnings"]

    try:
        presentation = Presentation(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as error:
        raise ValueError(f"PPTX 파일을 읽을 수 없어요 ({type(error).__name__}).") from error

    counts: dict[str, int] = {}
    source_index = 0
    pictures = 0
    charts = 0

    for slide_index, slide in enumerate(presentation.slides):
        number = slide_index + 1
        document["pages"].append(page_entry(number, PAGE_SOURCE))

        title = slide.shapes.title  # 제목 자리표시자는 위치와 상관없이 먼저 읽는다.
        ordered = ([title] if title is not None else []) + [
            shape for shape in _walk_shapes(slide.shapes, MSO_SHAPE_TYPE.GROUP)
            if title is None or shape.shape_id != title.shape_id
        ]
        for shape in ordered:
            if _shape_type(shape) == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
                continue
            if getattr(shape, "has_chart", False) and shape.has_chart:
                charts += 1
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                if add_table_block(document, counts, grid=_table_grid(shape.table), page=number,
                                   source_index=source_index, page_number_source=PAGE_SOURCE):
                    source_index += 1
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = _clean_lines(paragraph.text or "")
                    if add_paragraph_block(document, counts, text=text, page=number,
                                           source_index=source_index, page_number_source=PAGE_SOURCE):
                        source_index += 1

    if not document["pages"]:
        raise ValueError("슬라이드가 없는 PPTX예요.")
    if not document["blocks"]:
        raise ValueError("글자가 없는 PPTX는 지원하지 않아요. 글자가 포함된 PPTX를 올려 주세요.")

    warnings.append({"code": "pptx_images_not_extracted", "pictures": pictures, "charts": charts,
                     "detail": "그림·차트·슬라이드 노트는 추출하지 않습니다."})
    return finish_document(document)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("사용법: python pptx_parser.py <제안서.pptx>")
    result = parse_pptx(sys.argv[1])
    print("파싱 완료:", result["filename"])
    print("페이지:", len(result["pages"]), "블록:", len(result["blocks"]),
          "표:", len(result["tables"]), "경고:", len(result["warnings"]))
