"""제안서평가 집계표(정성) 엑셀 생성기.

공식 서식(집계표(제안서평가 서식) 샘플)을 openpyxl로 재현한다.

- "집계표" 시트: 사업명 제목, 위원별 업체 총점, 소계(A)·제외점수(B: 최고점/최저점)·
  합계(C=A-B)·최종평점, 확인 문구, 날짜, 서명란(평가위원장/검토자/작성자).
- 위원별 시트(시트명 = 위원 성명): 평가항목별 업체 점수표, 합계, 허용 점수 단계
  (배점 × 1.0/0.8/0.6/0.4), 오류검증 블록("배점에러" 표기).

모든 계산 셀은 서버가 서명된 정책으로 이미 계산한 값을 그대로 기록하며,
수식은 일절 쓰지 않는다. 업체 수(1~8)와 위원 수(3~11)에 따라 격자가 늘고 준다.
"""

from __future__ import annotations

import re
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ── 서식 상수 (샘플 통합 문서에서 추출) ─────────────────────────────────────
_FONT_TITLE = "경기천년제목 Medium"
_FONT_HEAD = "경기천년제목 Light"
_FONT_BODY = "경기천년바탕 Regular"
_FONT_BASIC = "맑은 고딕"

_FILL_HEADER = PatternFill("solid", fgColor="F2F2F2")  # 테마 흰색 -5%
_FILL_SUBTOTAL = PatternFill("solid", fgColor="DDEBF7")  # 강조5 +80%
_FILL_EXCLUDED = PatternFill("solid", fgColor="D9D9D9")  # 테마 흰색 -15%
_FILL_FINAL = PatternFill("solid", fgColor="FFF2CC")  # 강조4 +80%
_FILL_LEVELS = PatternFill("solid", fgColor="BFBFBF")  # 테마 흰색 -25%
_FILL_ERRORS = PatternFill("solid", fgColor="E2EFDA")  # 강조6 +80%

_THIN = Side(style="thin")
_MEDIUM = Side(style="medium")
_NONE = Side(style=None)

_CENTER = Alignment(horizontal="center", vertical="center")
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center")
_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
_VCENTER = Alignment(vertical="center")

_LEVEL_RATIOS = (1.0, 0.8, 0.6, 0.4)
_ERROR_TEXT = "배점에러"
_SUMMARY_SHEET = "집계표"
_INVALID_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")


# ── 공개 API ───────────────────────────────────────────────────────────────
def build_aggregation_workbook(report: dict) -> bytes:
    """집계 보고서(dict)로 집계표 통합 문서를 만들어 xlsx 바이트를 돌려준다."""
    vendors: list[dict] = list(report.get("vendors") or [])
    criteria: list[dict] = list(report.get("criteria") or [])
    evaluators: list[dict] = list(report.get("evaluators") or [])
    summary: dict[str, dict] = dict(report.get("vendor_summary") or {})
    signers: dict[str, Any] = dict(report.get("signers") or {})

    wb = Workbook()
    ws = wb.active
    ws.title = _SUMMARY_SHEET
    _write_summary_sheet(ws, report, vendors, evaluators, summary, signers)

    used_titles = {_SUMMARY_SHEET}
    for index, evaluator in enumerate(evaluators, start=1):
        title = _unique_sheet_title(str(evaluator.get("name") or f"위원{index}"), used_titles)
        used_titles.add(title)
        _write_evaluator_sheet(wb.create_sheet(title), evaluator, vendors, criteria)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ── 집계표 시트 ────────────────────────────────────────────────────────────
def _write_summary_sheet(
    ws: Worksheet,
    report: dict,
    vendors: list[dict],
    evaluators: list[dict],
    summary: dict[str, dict],
    signers: dict[str, Any],
) -> None:
    vendor_count = max(len(vendors), 1)
    evaluator_count = len(evaluators)
    col_no = 2  # B 구분
    col_name = 3  # C 위원성명
    col_v0 = 4  # D 첫 업체
    col_note = col_v0 + vendor_count  # 비고
    col_margin = col_note + 1  # 인쇄 여백 열

    row_header = 5
    row_e0 = 6
    row_sub = row_e0 + evaluator_count
    row_max = row_sub + 1
    row_min = row_sub + 2
    row_sum = row_sub + 3
    row_avg = row_sub + 4
    row_confirm = row_avg + 2
    row_date = row_avg + 4
    row_sign = row_avg + 6

    # 열 너비 / 행 높이
    ws.column_dimensions["A"].width = 1.75
    for col in range(col_no, col_note + 1):
        ws.column_dimensions[get_column_letter(col)].width = 11.125
    ws.column_dimensions[get_column_letter(col_margin)].width = 1.625
    heights = {2: 18.75, 3: 33.0, 4: 17.25, row_header: 36.0}
    for row in range(row_e0, row_avg):
        heights[row] = 25.5
    heights.update(
        {
            row_avg: 36.0,
            row_avg + 1: 17.25,
            row_confirm: 22.5,
            row_avg + 3: 17.25,
            row_date: 22.5,
            row_avg + 5: 16.5,
            row_sign: 23.1,
            row_sign + 1: 23.1,
            row_sign + 2: 23.1,
        }
    )
    for row, height in heights.items():
        ws.row_dimensions[row].height = height

    # 제목
    ws.merge_cells(start_row=2, start_column=col_no, end_row=3, end_column=col_note)
    title = ws.cell(row=2, column=col_no)
    title.value = f"【 {report.get('project_name', '')} 】\n제안서평가 집계표(정성)"
    title.font = Font(name=_FONT_TITLE, size=15, color="FF0000FF")
    title.alignment = _CENTER_WRAP

    head_font = Font(name=_FONT_HEAD, size=14, bold=True)
    body_font = Font(name=_FONT_HEAD, size=14, bold=False)

    def outer(col: int, row: int, top_row: int, bottom_row: int) -> Border:
        return Border(
            left=_MEDIUM if col == col_no else _THIN,
            right=_MEDIUM if col == col_note else _THIN,
            top=_MEDIUM if row == top_row else _THIN,
            bottom=_MEDIUM if row == bottom_row else _THIN,
        )

    # 머리글
    headers: dict[int, str] = {col_no: "구분", col_name: "위원성명", col_note: "비고"}
    for offset, vendor in enumerate(vendors):
        headers[col_v0 + offset] = str(vendor.get("label") or vendor.get("display_code") or "")
    for col in range(col_no, col_note + 1):
        cell = ws.cell(row=row_header, column=col, value=headers.get(col))
        cell.font = head_font
        cell.fill = _FILL_HEADER
        cell.alignment = _CENTER_WRAP
        cell.border = outer(col, row_header, row_header, row_avg)

    # 위원별 행
    for offset, evaluator in enumerate(evaluators):
        row = row_e0 + offset
        scores = evaluator.get("scores") or {}
        for col in range(col_no, col_note + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = _CENTER
            cell.border = outer(col, row, row_header, row_avg)
            cell.font = head_font if col in (col_no, col_name) else body_font
        ws.cell(row=row, column=col_no, value=offset + 1)
        ws.cell(row=row, column=col_name, value=evaluator.get("name"))
        for v_offset, vendor in enumerate(vendors):
            vendor_scores = scores.get(vendor.get("display_code")) or {}
            ws.cell(row=row, column=col_v0 + v_offset, value=_num(vendor_scores.get("total")))
        if evaluator.get("is_chair"):
            ws.cell(row=row, column=col_note, value="위원장")

    # 소계 / 제외점수 / 합계 / 최종평점
    included_count = _included_count(summary, evaluator_count)
    summary_rows: list[tuple[int, str, PatternFill, str, str | None]] = [
        (row_sub, "소계(A)", _FILL_SUBTOTAL, "subtotal", None),
        (row_max, "제외점수(B)", _FILL_EXCLUDED, "max", "최고점"),
        (row_min, "", _FILL_EXCLUDED, "min", "최저점"),
        (row_sum, "합계(C=A-B)", _FILL_SUBTOTAL, "included_sum", None),
        (row_avg, f"최종평점\n(평점=C/{included_count})", _FILL_FINAL, "average", None),
    ]
    for row, label, fill, key, note in summary_rows:
        for col in range(col_no, col_note + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = outer(col, row, row_header, row_avg)
            cell.alignment = _CENTER_WRAP if col == col_no else _CENTER
            cell.font = head_font if col == col_no else body_font
            if col < col_note:
                cell.fill = fill
        if row != row_min:
            ws.cell(row=row, column=col_no, value=label)
        for v_offset, vendor in enumerate(vendors):
            cell = ws.cell(row=row, column=col_v0 + v_offset)
            vendor_summary = summary.get(vendor.get("display_code")) or {}
            raw = vendor_summary.get(key)
            if key == "average":
                cell.value = _average(raw)
                cell.number_format = "0.00"
            else:
                cell.value = _num(raw)
        if note:
            ws.cell(row=row, column=col_note, value=note)
    ws.merge_cells(start_row=row_sub, start_column=col_no, end_row=row_sub, end_column=col_name)
    ws.merge_cells(start_row=row_max, start_column=col_no, end_row=row_min, end_column=col_name)
    ws.merge_cells(start_row=row_sum, start_column=col_no, end_row=row_sum, end_column=col_name)
    ws.merge_cells(start_row=row_avg, start_column=col_no, end_row=row_avg, end_column=col_name)

    # 확인 문구 / 날짜 / 서명
    ws.merge_cells(start_row=row_confirm, start_column=col_no, end_row=row_confirm, end_column=col_note)
    confirm = ws.cell(row=row_confirm, column=col_no)
    confirm.value = "위와 같은 결과로 공정하게 평가하였음을 확인합니다."
    confirm.font = Font(name=_FONT_BODY, size=18, bold=True)
    confirm.alignment = _CENTER

    date_end = max(col_v0, col_note - 2)
    if date_end > col_v0:
        ws.merge_cells(start_row=row_date, start_column=col_v0, end_row=row_date, end_column=date_end)
    date_cell = ws.cell(row=row_date, column=col_v0)
    date_cell.value = _korean_date(str(report.get("finalized_at") or ""))
    date_cell.number_format = "@"
    date_cell.font = Font(name=_FONT_BODY, size=18)
    date_cell.alignment = _CENTER

    sign_font = Font(name=_FONT_BODY, size=18)
    sign_lines = (
        f" 평가위원장   {signers.get('chair') or ''} (서명) ",
        f" 검  토  자   {signers.get('reviewer') or ''} (서명)",
        f" 작  성  자   {signers.get('author') or ''} (서명)",
    )
    sign_end = max(col_name, col_note - 1)
    for offset, line in enumerate(sign_lines):
        row = row_sign + offset
        ws.merge_cells(start_row=row, start_column=col_name, end_row=row, end_column=sign_end)
        cell = ws.cell(row=row, column=col_name, value=line)
        cell.font = sign_font
        cell.alignment = _LEFT

    # 인쇄 설정 (A4 세로)
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "portrait"
    ws.print_area = f"A1:{get_column_letter(col_margin)}{row_sign + 3}"


# ── 위원별 시트 ────────────────────────────────────────────────────────────
def _write_evaluator_sheet(
    ws: Worksheet,
    evaluator: dict,
    vendors: list[dict],
    criteria: list[dict],
) -> None:
    vendor_count = max(len(vendors), 1)
    criteria_count = len(criteria)
    col_label = 2  # B 구분
    col_v0 = 3  # C 첫 업체
    col_note = col_v0 + vendor_count  # 비고
    col_l0 = col_note + 2  # 허용 점수(4열)
    col_e0 = col_l0 + len(_LEVEL_RATIOS) + 1  # 오류검증
    col_e_end = col_e0 + vendor_count - 1

    row_title = 2
    row_band = 3
    row_header = 4
    row_c0 = 5
    row_total = row_c0 + criteria_count

    ws.column_dimensions["B"].width = 31.5
    for col in range(col_v0, col_note):
        ws.column_dimensions[get_column_letter(col)].width = 10.0
    for col in range(col_l0, col_l0 + len(_LEVEL_RATIOS)):
        ws.column_dimensions[get_column_letter(col)].width = 5.0
    ws.row_dimensions[row_title].height = 33.75
    ws.row_dimensions[row_band].height = 20.25
    for row in range(row_header, row_total + 1):
        ws.row_dimensions[row].height = 39.95

    box = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
    head_font = Font(name=_FONT_HEAD, size=14, color="FF000000")
    score_font = Font(name=_FONT_BASIC, size=14)
    level_font = Font(name=_FONT_BASIC, size=10)
    error_font = Font(name=_FONT_BASIC, size=11, color="FFFF0000")
    band_font = Font(name=_FONT_BASIC, size=11)

    # 제목 / 보조 블록 머리띠
    ws.merge_cells(start_row=row_title, start_column=col_label, end_row=row_title, end_column=col_note)
    title = ws.cell(row=row_title, column=col_label, value=f"{evaluator.get('name') or ''} 위원")
    title.font = Font(name=_FONT_BASIC, size=22)
    title.fill = _FILL_SUBTOTAL
    title.alignment = _CENTER

    ws.merge_cells(
        start_row=row_band, start_column=col_l0, end_row=row_band, end_column=col_l0 + len(_LEVEL_RATIOS) - 1
    )
    levels_band = ws.cell(row=row_band, column=col_l0, value="허용 점수")
    levels_band.font = band_font
    levels_band.fill = _FILL_LEVELS
    levels_band.alignment = _CENTER

    ws.merge_cells(start_row=row_band, start_column=col_e0, end_row=row_band, end_column=col_e_end)
    errors_band = ws.cell(row=row_band, column=col_e0, value="오류검증탭")
    errors_band.font = band_font
    errors_band.fill = _FILL_ERRORS
    errors_band.alignment = _CENTER

    # 머리글
    headers: dict[int, str] = {col_label: "구분", col_note: "비고"}
    for offset, vendor in enumerate(vendors):
        headers[col_v0 + offset] = str(vendor.get("label") or vendor.get("display_code") or "")
    for col in range(col_label, col_note + 1):
        cell = ws.cell(row=row_header, column=col, value=headers.get(col))
        cell.font = head_font
        cell.fill = _FILL_HEADER
        cell.alignment = _CENTER_WRAP
        cell.border = box

    # 평가항목 행
    scores: dict[str, dict] = evaluator.get("scores") or {}
    for offset, criterion in enumerate(criteria):
        row = row_c0 + offset
        code = str(criterion.get("code") or "")
        max_score = float(criterion.get("max_score") or 0)
        levels = [round(max_score * ratio, 4) for ratio in _LEVEL_RATIOS]

        label = ws.cell(row=row, column=col_label, value=f"{code}. {criterion.get('title') or ''}")
        label.font = head_font
        label.fill = _FILL_HEADER
        label.alignment = _LEFT_WRAP
        label.border = box

        for v_offset, vendor in enumerate(vendors):
            vendor_scores = scores.get(vendor.get("display_code")) or {}
            score = (vendor_scores.get("items") or {}).get(code)
            cell = ws.cell(row=row, column=col_v0 + v_offset, value=_num(score))
            cell.font = score_font
            cell.alignment = _CENTER
            cell.border = box

            error = ws.cell(row=row, column=col_e0 + v_offset)
            error.value = None if _is_allowed(score, levels) else _ERROR_TEXT
            error.font = error_font
            error.fill = _FILL_ERRORS
            error.alignment = _VCENTER

        note = ws.cell(row=row, column=col_note)
        note.font = score_font
        note.alignment = _VCENTER
        note.border = box

        for l_offset, level in enumerate(levels):
            cell = ws.cell(row=row, column=col_l0 + l_offset, value=_num(level))
            cell.font = level_font
            cell.fill = _FILL_LEVELS
            cell.alignment = _CENTER
            cell.border = box

    # 합계 행
    total_label = ws.cell(row=row_total, column=col_label, value="합계")
    total_label.font = head_font
    total_label.fill = _FILL_FINAL
    total_label.alignment = _CENTER_WRAP
    total_label.border = box
    for v_offset, vendor in enumerate(vendors):
        vendor_scores = scores.get(vendor.get("display_code")) or {}
        cell = ws.cell(row=row_total, column=col_v0 + v_offset, value=_num(vendor_scores.get("total")))
        cell.font = score_font
        cell.fill = _FILL_FINAL
        cell.alignment = _CENTER
        cell.border = box
    note = ws.cell(row=row_total, column=col_note)
    note.font = score_font
    note.alignment = _VCENTER
    note.border = box
    for l_offset in range(len(_LEVEL_RATIOS)):
        cell = ws.cell(row=row_total, column=col_l0 + l_offset)
        cell.font = level_font
        cell.fill = _FILL_LEVELS
        cell.alignment = _VCENTER
        cell.border = box


# ── 보조 함수 ──────────────────────────────────────────────────────────────
def _num(value: Any) -> int | float | None:
    """정수로 떨어지는 실수는 정수로, 그 외는 그대로. 숫자가 아니면 None."""
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else number


def _average(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _is_allowed(score: Any, levels: list[float]) -> bool:
    number = _num(score)
    if number is None:
        return False
    return any(abs(float(number) - level) < 1e-6 for level in levels)


def _included_count(summary: dict[str, dict], evaluator_count: int) -> int:
    for vendor_summary in summary.values():
        count = (vendor_summary or {}).get("included_count")
        if isinstance(count, int) and count > 0:
            return count
    return max(evaluator_count - 2, 1)


def _korean_date(raw: str) -> str:
    """'YYYY-MM-DD'(또는 ISO datetime) → 'YYYY년 M월 D일'. 실패 시 원문."""
    text = raw.strip()
    try:
        parsed = date.fromisoformat(text[:10])
    except ValueError:
        return text
    return f"{parsed.year}년 {parsed.month}월 {parsed.day}일"


def _unique_sheet_title(name: str, used: set[str]) -> str:
    base = _INVALID_TITLE_CHARS.sub("", name).strip() or "위원"
    base = base[:31]
    if base not in used:
        return base
    for n in range(2, 1000):
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        if candidate not in used:
            return candidate
    raise ValueError("시트 이름을 만들 수 없습니다")
