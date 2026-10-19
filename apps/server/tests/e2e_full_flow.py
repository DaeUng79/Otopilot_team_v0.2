r"""전 과정 HTTP 자동 시험 (2026-10-18). 실제 데모 DB가 아니라 복사본으로 띄운 서버에 대고 실행한다.

준비:
  1) 데모 DB·문서 DB를 임시 폴더에 복사 (admin.sqlite3, evaluation.sqlite3, keys.bin, access-code.txt, documents_parse.sqlite3)
  2) .venv\Scripts\python.exe apps/server/run.py --demo --port 5199 --data-dir <임시>\data --documents-db <임시>\documents_parse.sqlite3
  3) .venv\Scripts\python.exe -I apps/server/tests/e2e_full_flow.py <결과 저장 폴더> "C:\coding\제안요청서, 제안서 샘플"
전제: seed_samples.py가 넣은 sample-medical-ai(평가표 확정·업체 A·B 제안서 READY)와 sample-mydata-2026이 있어야 한다.
마지막 줄 "FAILURES: none"이면 통과.
"""
import base64, io, json, struct, sys, time, uuid, zlib, urllib.request, urllib.error
from http.cookiejar import CookieJar
from pathlib import Path
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
BASE = "http://127.0.0.1:5199"
OUT = Path(sys.argv[1])
SAMPLES = Path(sys.argv[2])
PID = "sample-medical-ai"
failures = []


def png(w=200, h=100):
    raw = b"".join(b"\x00" + bytes([0, 0, 0, 255] * w) for _ in range(h))
    def chunk(t, d): return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 6, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


class Actor:
    def __init__(self, label):
        self.label = label
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(CookieJar()))

    def call(self, method, path, body=None, raw=None, headers=None, expect=None):
        data = None
        h = {"X-Simsa-Client": "v2-web", "Idempotency-Key": str(uuid.uuid4())}
        if body is not None:
            data = json.dumps(body).encode(); h["Content-Type"] = "application/json"
        if raw is not None:
            data, ctype = raw; h["Content-Type"] = ctype
        if headers: h.update(headers)
        req = urllib.request.Request(BASE + path, data=data, method=method, headers=h)
        try:
            with self.opener.open(req, timeout=120) as r:
                content = r.read(); status = r.status; ctype = r.headers.get("content-type", "")
        except urllib.error.HTTPError as e:
            content = e.read(); status = e.code; ctype = e.headers.get("content-type", "")
        if expect is not None and status != expect:
            raise AssertionError(f"{method} {path} -> {status} (expected {expect}): {content[:300]!r}")
        if "json" in ctype:
            return status, json.loads(content)
        return status, content

    def act(self, action, body, expect=200):
        status, data = self.call("POST", f"/api/v2/actions/{action}", body, expect=expect)
        return data.get("data", data) if isinstance(data, dict) else data


def step(name, fn):
    try:
        result = fn(); print(f"PASS {name}" + (f" · {result}" if result else "")); return result
    except Exception as exc:
        failures.append(name); print(f"FAIL {name} · {type(exc).__name__}: {exc}"); return None


def multipart(fields, filename, data):
    boundary = "----e2e" + uuid.uuid4().hex
    body = b""
    for k, v in fields.items():
        body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode()
    body += f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\nContent-Type: application/octet-stream\r\n\r\n".encode() + data + f"\r\n--{boundary}--\r\n".encode()
    return body, f"multipart/form-data; boundary={boundary}"


for _ in range(40):
    try:
        urllib.request.urlopen(BASE + "/api/v2/config", timeout=3); break
    except Exception:
        time.sleep(1)
admin = Actor("admin")
step("관리자 로그인", lambda: admin.act("AUTH-01", {"username": "web-demo", "password": "DemoOnly!2026"})["role"])
pages = ["dept-admin/PM-02/design/documents.html", "dept-admin/PM-03/design/criteria.html", "dept-admin/PM-04/design/committee.html",
         "dept-admin/PM-05/design/attendance.html", "dept-admin/PM-06/design/results.html", "dept-admin/PM-07/design/pool.html",
         "super-admin/SA-03/design/projects.html", "super-admin/SA-04/design/data.html"]
for p in pages:
    step(f"화면 렌더 {p.split('/')[1]}", lambda p=p: admin.call("GET", "/app/" + p, expect=200)[0])
state = step("평가표 상태", lambda: admin.call("GET", f"/api/v2/projects/{PID}/criteria", expect=200)[1]["data"])
assert state["versions"][-1]["status"] == "APPROVED" and state["versions"][-1]["ack"], state["versions"][-1]
detail = step("사업 상세", lambda: admin.act("PRJ-05", {"project_id": PID}))
vendor_c = next(v for v in detail["vendors"] if v["display_code"] == "C")
step("업체 C 삭제(제안서 없음)", lambda: admin.act("PRJ-07", {"project_id": PID, "vendor_id": vendor_c["id"], "version": vendor_c["version"]}))
step("정원 4명 거부", lambda: admin.call("POST", "/api/v2/actions/DRAW-01", {"project_id": PID, "selection_mode": "POOL_RANDOM", "target_count": 4, "reserve_count": 2, "quotas": [{"specialty_id": "AI", "final_count": 4}], "version": 0}, expect=422)[0])
rule = step("규칙 저장 정원 5", lambda: admin.act("DRAW-01", {"project_id": PID, "selection_mode": "POOL_RANDOM", "target_count": 5, "reserve_count": 2, "quotas": [{"specialty_id": "AI", "final_count": 3}, {"specialty_id": "MED", "final_count": 1}, {"specialty_id": "DATA", "final_count": 1}], "version": 0}))
draw = step("3배수 추첨", lambda: admin.act("DRAW-02", {"project_id": PID, "rule_version": rule["version"]}))
assert len(draw["members"]) == 15, len(draw["members"])
committee = admin.call("GET", f"/committee-v2/api/projects/{PID}", expect=200)[1]["data"]
items = committee["items"]
by_field = {}
for n in items:
    f = next(d["specialty_id"] for d in committee["draw_members"] if d["id"] == n["draw_member_id"])
    by_field.setdefault(f, []).append(n)
chosen = by_field["AI"][:3] + by_field["MED"][:1] + by_field["DATA"][:1]
reserves = by_field["AI"][3:5]
for n in chosen + reserves:
    r = admin.act("ROS-03", {"project_id": PID, "nominee_id": n["id"], "version": n["version"], "contact_status": "AVAILABLE", "contact_note": "e2e"})
    n["version"] = r["version"]
members = [{"nominee_id": n["id"], "version": n["version"], "member_kind": "PRIMARY", "reserve_order": None} for n in chosen]
step("예비 0명 확정 거부", lambda: admin.call("POST", "/api/v2/actions/ROS-05", {"project_id": PID, "rule_version": rule["version"], "members": members, "reserve_count": 0}, expect=422)[0])
roster = step("명단 확정 5+2", lambda: admin.act("ROS-05", {"project_id": PID, "rule_version": rule["version"], "members": members + [{"nominee_id": r["id"], "version": r["version"], "member_kind": "RESERVE", "reserve_order": i + 1} for i, r in enumerate(reserves)], "reserve_count": 2}))
step("확정명단 xlsx", lambda: len(admin.call("POST", "/api/v2/actions/ROS-06", {"project_id": PID, "roster_id": roster["id"]}, expect=200)[1]))
evaluators = []
for n in chosen:
    ev = Actor(n["candidate"]["name"])
    step(f"위원 로그인 {ev.label}", lambda ev=ev, n=n: ev.act("AUTH-02", {"name": n["candidate"]["name"], "last4": n["candidate"]["phone"][-4:], "project_id": PID})["role"])
    step(f"참석 확인 {ev.label}", lambda ev=ev: ev.act("ATT-01", {"status": "CONFIRMED", "consent": True, "consent_version": "attendance-v2", "version": 1})["status"])
    evaluators.append(ev)
att = step("참석 현황", lambda: admin.act("ATT-02", {"project_id": PID}))
session = step("회차 준비 정원5", lambda: admin.act("ATT-00", {"project_id": PID, "roster_id": att["rosters"][-1]["id"], "approved_criteria_version_id": att["criteria"][-1]["id"], "target_count": 5, "minimum_attendees": 3, "approve_policy": True}))
att = admin.act("ATT-02", {"project_id": PID})
ids = [a["attendance_id"] for a in att["assignments"] if a["attendance_status"] == "CONFIRMED" and a["member_kind"] == "PRIMARY"]
session = step("실제 평가인원 확정", lambda: admin.act("ATT-03", {"session_id": session["id"], "version": session["version"], "attendance_ids": ids}))
att = admin.act("ATT-02", {"project_id": PID})
chair_participant = att["participants"][0]
session = step("위원장 지정", lambda: admin.act("ATT-04", {"session_id": session["id"], "version": session["version"], "participant_id": chair_participant["id"]}))
session = step("평가 시작", lambda: admin.act("ATT-05", {"session_id": session["id"], "version": session["version"]}))
sid = session["id"]
chair = None
for k, ev in enumerate(evaluators):
    sheets = ev.act("SCO-01", {"session_id": sid})
    if sheets["participant"]["id"] == chair_participant["id"]:
        chair = ev
    for j, sh in enumerate(sheets["sheets"]):
        entries = [{"criterion_id": e["criterion_id"], "score": max(0, e["max_score_snapshot"] - ((k + j + i) % 3) * 10000), "comment": None} for i, e in enumerate(sh["entries"])]
        ev.act("SCO-02", {"session_id": sid, "sheet_id": sh["id"], "version": sh["version"], "opinion": "e2e 의견", "entries": entries})
    review = ev.act("SCO-03", {"session_id": sid})
    sig = ev.act("SCO-04", {"session_id": sid, "purpose": "EVALUATION", "payload_sha256": review["payload_sha256"], "image": "data:image/png;base64," + base64.b64encode(png()).decode(), "consent": True, "consent_text_version": "signature-v2"})
    step(f"제출 {ev.label}", lambda ev=ev, sig=sig, review=review: ev.act("SCO-05", {"session_id": sid, "signature_id": sig["id"], "payload_sha256": review["payload_sha256"]})["submitted_count"])
run = step("위원장 집계", lambda: chair.act("RES-02", {"session_id": sid}))
res = chair.act("RES-01", {"session_id": sid})
assert res["is_chair"] and res["run"]["id"] == run["run"]["id"]
exclusions = [i["exclusion"] for i in res["inputs"]]
assert exclusions.count("LOWEST") == len(res["items"]) and exclusions.count("HIGHEST") == len(res["items"]), exclusions
csig = step("위원장 서명", lambda: chair.act("RES-03", {"session_id": sid, "purpose": "CHAIR_RESULT", "result_run_id": run["run"]["id"], "payload_sha256": res["payload_sha256"], "image": "data:image/png;base64," + base64.b64encode(png()).decode(), "consent": True, "consent_text_version": "signature-v2"}))
step("최종 결과 제출", lambda: chair.act("RES-04", {"session_id": sid, "result_run_id": run["run"]["id"], "signature_id": csig["id"], "payload_sha256": res["payload_sha256"]})["status"])
final = admin.act("RES-01", {"session_id": sid})
assert final["session"]["status"] == "FINALIZED", final["session"]["status"]
xlsx = step("집계표 xlsx 다운로드", lambda: admin.call("POST", "/api/v2/actions/RES-05", {"session_id": sid, "format": "xlsx"}, expect=200)[1])
(OUT / "집계표.xlsx").write_bytes(xlsx)
html = step("서명증빙 html 다운로드", lambda: admin.call("POST", "/api/v2/actions/RES-05", {"session_id": sid}, expect=200)[1])
(OUT / "결과.html").write_bytes(html)
def check_xlsx():
    from openpyxl import load_workbook
    wb = load_workbook(OUT / "집계표.xlsx"); ws = wb["집계표"]
    rows = [[c for c in r if c is not None] for r in ws.iter_rows(min_row=5, max_row=16, values_only=True)]
    return wb.sheetnames, rows[:9]
step("집계표 내용", check_xlsx)
step("사업 상태 COMPLETED", lambda: admin.act("PRJ-05", {"project_id": PID})["project"]["status"])
mats = step("자료 목록", lambda: admin.call("GET", "/api/v2/materials", expect=200)[1]["data"])
assert any(r["session_id"] == sid for r in mats["results"]) and any(i["kind"] == "PROPOSAL" for i in mats["items"])
body, ctype = multipart({"title": "심사 회의록", "description": "e2e"}, "회의록.html", b"<html>e2e</html>")
mat = step("결과 자료 등록 MAT-02", lambda: admin.call("POST", f"/api/v2/projects/{PID}/result-files", raw=(body, ctype), expect=201)[1]["data"])
body, ctype = multipart({"version": str(mat["version"]), "description": "수정본"}, "회의록_v2.html", b"<html>e2e v2</html>")
step("결과 자료 새 버전 MAT-03", lambda: admin.call("POST", f"/api/v2/projects/{PID}/result-files/{mat['id']}/revisions", raw=(body, ctype), expect=200)[1]["data"]["revision_no"])
# 다른 사업으로 제안서 등록·평가표 승인 HTTP 경로
P2 = "sample-mydata-2026"
d2 = admin.act("PRJ-05", {"project_id": P2})
vendor = d2["vendors"][0]
hwpx = next(SAMPLES.rglob("마이허브*.hwpx"))
body, ctype = multipart({}, hwpx.name, hwpx.read_bytes())
up = step("제안서 업로드 PRO-01 (HWPX)", lambda: admin.call("POST", f"/api/v2/projects/{P2}/vendors/{vendor['id']}/proposals", raw=(body, ctype), expect=201)[1]["data"])
assert up["ack"]["status"] == "APPLIED", up["ack"]
step("제안서 목록 PRO-02", lambda: len(admin.call("GET", f"/api/v2/projects/{P2}/proposals", expect=200)[1]["data"]["uploads"]))
step("제안서 원본 다운로드", lambda: len(admin.call("GET", f"/api/v2/projects/{P2}/proposals/{up['id']}/download", expect=200)[1]))
step("제안서 파싱 원문", lambda: len(admin.call("GET", f"/api/v2/projects/{P2}/proposals/{up['id']}/source", expect=200)[1]["data"]["blocks"]))
step("잘못된 형식 거부", lambda: admin.call("POST", f"/api/v2/projects/{P2}/vendors/{vendor['id']}/proposals", raw=multipart({}, "x.txt", b"hello"), expect=422)[0])
draft = step("평가표 서식 복사본 CRT-03", lambda: admin.call("POST", f"/api/v2/projects/{P2}/criteria-drafts", {"source": "template"}, expect=201)[1]["data"])
items = [{"code": i["code"], "parent_code": i["parent_code"], "title": i["title"], "description": i["description"], "max_score": i["points"], "source_page": i["source_page"] or 3, "sort_order": i["sort_order"]} for i in draft["items"]]
items[0]["title"] = "사업 이해도 및 추진 전략"
step("배점 합 불일치 거부 CRT-04", lambda: admin.call("PUT", f"/api/v2/projects/{P2}/criteria-drafts/{draft['id']}", {"version": draft["version"], "items": items[:-1], "total_points": draft["total_points"]}, expect=422)[0])
saved = step("평가표 저장 CRT-04", lambda: admin.call("PUT", f"/api/v2/projects/{P2}/criteria-drafts/{draft['id']}", {"version": draft["version"], "items": items, "total_points": draft["total_points"]}, expect=200)[1]["data"])
approved = step("평가표 확정 CRT-05", lambda: admin.call("POST", f"/api/v2/projects/{P2}/criteria-drafts/{draft['id']}/approve", {"version": saved["version"], "confirmed": True}, expect=200)[1]["data"])
assert approved["ack"]["status"] == "APPLIED", approved["ack"]
step("제안서 보관 처리 PRO-03", lambda: admin.call("DELETE", f"/api/v2/projects/{P2}/proposals/{up['id']}", {"reason": "e2e 테스트"}, expect=200)[1]["data"]["status"])
step("요구사항 상태 DOC-03", lambda: admin.call("GET", f"/api/v2/projects/{P2}/requirements", expect=200)[1]["data"]["status"])
print("\nFAILURES:", failures or "none")
